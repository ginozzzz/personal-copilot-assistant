from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


@dataclass
class StrikeHit:
    file_path: Path
    sheet_name: str
    row_number: int
    cell_coordinate: str
    value: object


def iter_xlsx_files(root: Path) -> Iterable[Path]:
    for path in sorted(root.rglob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        yield path


def find_strikethrough_hits(file_path: Path) -> list[StrikeHit]:
    workbook = load_workbook(filename=file_path, read_only=False, data_only=False)
    hits: list[StrikeHit] = []

    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.font and cell.font.strike:
                        hits.append(
                            StrikeHit(
                                file_path=file_path,
                                sheet_name=sheet.title,
                                row_number=cell.row,
                                cell_coordinate=cell.coordinate,
                                value=cell.value,
                            )
                        )
    finally:
        workbook.close()

    return hits


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Recursively scan XLSX files and report cells using strikethrough formatting."
    )
    parser.add_argument("folder", type=Path, help="Folder to scan recursively")
    args = parser.parse_args()

    root = args.folder.expanduser().resolve()
    if not root.exists() or not root.is_dir():
        print(f"ERROR\tFolder not found or not a directory\t{root}")
        return 2

    files = list(iter_xlsx_files(root))
    if not files:
        print(f"SCANNED_FILES\t0")
        print("NO_XLSX_FILES_FOUND")
        return 0

    all_hits: list[StrikeHit] = []
    files_with_hits: set[Path] = set()

    for file_path in files:
        hits = find_strikethrough_hits(file_path)
        if hits:
            files_with_hits.add(file_path)
            all_hits.extend(hits)

    print(f"SCANNED_FILES\t{len(files)}")
    print(f"FILES_WITH_STRIKETHROUGH\t{len(files_with_hits)}")
    print(f"STRIKETHROUGH_CELLS\t{len(all_hits)}")

    if not all_hits:
        print("NO_STRIKETHROUGHS_FOUND")
        return 0

    print("FILE\tSHEET\tROW\tCELL\tVALUE")
    for hit in all_hits:
        value = "" if hit.value is None else str(hit.value).replace("\n", " ")
        print(
            f"{hit.file_path}\t{hit.sheet_name}\t{hit.row_number}\t{hit.cell_coordinate}\t{value}"
        )

    return 0


if __name__ == "__main__":
    raise SystemExit(main())